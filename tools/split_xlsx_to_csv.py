#!/usr/bin/env python3
"""
Split XLSX to CSV
=================
Tool untuk membaca file Excel (.xlsx), memilih kolom yang dibutuhkan,
lalu split menjadi multiple file CSV berdasarkan jumlah baris per file.

Cara pakai: python split_xlsx_to_csv.py
Script akan interaktif menanyakan:
  1. Path file .xlsx
  2. Kolom mana saja yang dibutuhkan
  3. Berapa baris per 1 file CSV output

Author: Hermes Agent for Rholly
"""

import os
import sys
import csv
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: Package 'openpyxl' belum terpasang.")
    print("Install dengan: pip install openpyxl")
    sys.exit(1)


# ─── Helper: warna CLI ───
class c:
    HEADER = '\033[95m'
    OK = '\033[92m'
    WARN = '\033[93m'
    ERR = '\033[91m'
    CYAN = '\033[96m'
    BOLD = '\033[1m'
    END = '\033[0m'


def banner():
    print(f"""
{c.CYAN}{c.BOLD}╔══════════════════════════════════════════╗
║   Split XLSX to CSV — Interactive Tool   ║
╚══════════════════════════════════════════╝{c.END}
""")


def ask_file_path() -> str:
    """Tanya path file xlsx ke user, validasi sampai benar."""
    while True:
        path = input(f"{c.BOLD}📂 Masukkan path file .xlsx: {c.END}").strip()
        if not path:
            print(f"{c.WARN}  ⚠ Path tidak boleh kosong.{c.END}")
            continue
        if not os.path.isfile(path):
            print(f"{c.ERR}  ❌ File tidak ditemukan: {path}{c.END}")
            continue
        if not path.lower().endswith(('.xlsx', '.xlsm')):
            print(f"{c.WARN}  ⚠ File bukan .xlsx. Yakin lanjut? (y/n): {c.END}", end=' ')
            if input().strip().lower() != 'y':
                continue
        return os.path.abspath(path)


def read_headers(filepath: str) -> tuple:
    """Baca file xlsx, return (workbook, worksheet, headers)."""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    print(f"\n{c.OK}  ✓ Sheet: {ws.title} | Dimensi: {ws.dimensions}{c.END}")

    headers = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(cell) if cell is not None else f"Col_{i+1}" for i, cell in enumerate(row)]
        break

    if not headers:
        print(f"{c.ERR}  ❌ File tidak punya header (kosong?).{c.END}")
        sys.exit(1)

    return wb, ws, headers


def show_columns(headers: list) -> list:
    """Tampilkan semua kolom, biarkan user pilih mana yang dipakai."""
    print(f"\n{c.BOLD}📋 Kolom yang tersedia:{c.END}\n")
    for i, h in enumerate(headers, 1):
        print(f"  {c.CYAN}{i:>3}.{c.END} {h}")
    print()

    while True:
        inp = input(
            f"{c.BOLD}Pilih kolom (pisahkan dengan koma, contoh: 1,3,5 "
            f"atau ketik 'all' untuk semua): {c.END}"
        ).strip()

        if inp.lower() == 'all':
            return list(range(len(headers)))

        try:
            indices = [int(x.strip()) - 1 for x in inp.split(',')]
            if all(0 <= i < len(headers) for i in indices):
                selected = [(i, headers[i]) for i in indices]
                print(f"\n{c.OK}  ✓ Kolom terpilih:{c.END}")
                for i, name in selected:
                    print(f"    {c.CYAN}{i+1}.{c.END} {name}")
                print()
                return indices
            else:
                print(f"{c.ERR}  ❌ Nomor kolom di luar range.{c.END}")
        except ValueError:
            print(f"{c.ERR}  ❌ Format salah. Gunakan angka dipisah koma, contoh: 1,3,5{c.END}")


def ask_split_size(total_rows: int) -> int:
    """Tanya berapa baris per file CSV."""
    print(f"{c.BOLD}📊 Total baris data: {total_rows}{c.END}\n")
    while True:
        inp = input(
            f"{c.BOLD}Split menjadi berapa baris per file? "
            f"(contoh: 1000, atau ketik 'all' untuk 1 file): {c.END}"
        ).strip()

        if inp.lower() == 'all':
            return total_rows

        try:
            size = int(inp)
            if size <= 0:
                print(f"{c.ERR}  ❌ Harus angka positif.{c.END}")
                continue
            files_count = (total_rows + size - 1) // size
            print(f"\n{c.CYAN}  → Akan menghasilkan {files_count} file CSV "
                  f"({size} baris/file){c.END}\n")
            confirm = input(f"{c.BOLD}  Lanjut? (y/n): {c.END}").strip().lower()
            if confirm == 'y':
                return size
        except ValueError:
            print(f"{c.ERR}  ❌ Masukkan angka yang valid.{c.END}")


def ask_output_dir(filepath: str) -> str:
    """Tanya direktori output untuk file CSV."""
    default_dir = os.path.join(os.path.dirname(filepath), "output_csv")
    inp = input(
        f"\n{c.BOLD}📁 Direktori output "
        f"(Enter untuk default: {default_dir}): {c.END}"
    ).strip()
    out_dir = inp if inp else default_dir
    os.makedirs(out_dir, exist_ok=True)
    print(f"{c.OK}  ✓ Output: {out_dir}{c.END}\n")
    return out_dir


def split_and_export(ws, headers, selected_indices, split_size, out_dir, base_name):
    """Baca data, split, dan export ke CSV."""
    total_written = 0
    file_num = 1
    batch = []
    selected_headers = [headers[i] for i in selected_indices]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    print(f"{c.CYAN}⏳ Memproses data...{c.END}")

    for row in ws.iter_rows(min_row=2, values_only=True):
        filtered = [str(row[i]) if i < len(row) and row[i] is not None else '' for i in selected_indices]
        batch.append(filtered)

        if len(batch) >= split_size:
            fname = f"{base_name}_part{file_num:03d}_{timestamp}.csv"
            fpath = os.path.join(out_dir, fname)
            _write_csv(fpath, selected_headers, batch)
            total_written += len(batch)
            print(f"  {c.OK}✓ {fname} ({len(batch)} rows){c.END}")
            batch = []
            file_num += 1

    # Sisa baris terakhir
    if batch:
        fname = f"{base_name}_part{file_num:03d}_{timestamp}.csv"
        fpath = os.path.join(out_dir, fname)
        _write_csv(fpath, selected_headers, batch)
        total_written += len(batch)
        print(f"  {c.OK}✓ {fname} ({len(batch)} rows){c.END}")

    return total_written, file_num


def _write_csv(filepath, headers, rows):
    """Tulis satu file CSV dengan BOM (utf-8-sig) agar Excel bisa baca."""
    with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def main():
    banner()

    # 1. Tanya path file
    filepath = ask_file_path()
    base_name = os.path.splitext(os.path.basename(filepath))[0]

    # 2. Baca xlsx dan tampilkan kolom
    wb, ws, headers = read_headers(filepath)

    # 3. Tanya kolom mana yang dibutuhkan
    selected_indices = show_columns(headers)

    # 4. Hitung total baris data
    total_rows = ws.max_row - 1
    if total_rows <= 0:
        print(f"{c.ERR}  ❌ Tidak ada baris data.{c.END}")
        sys.exit(1)

    # 5. Tanya split size
    split_size = ask_split_size(total_rows)

    # 6. Tanya direktori output
    out_dir = ask_output_dir(filepath)

    # 7. Proses split & export
    print(f"\n{c.BOLD}🚀 Mulai proses split...{c.END}\n")
    total_written, file_count = split_and_export(
        ws, headers, selected_indices, split_size, out_dir, base_name
    )

    # 8. Summary
    print(f"\n{c.OK}{c.BOLD}══════════════════════════════════════")
    print(f"  ✅ SELESAI!")
    print(f"  📊 Total baris: {total_written}")
    print(f"  📁 Total file:  {file_count}")
    print(f"  📂 Lokasi:      {out_dir}")
    print(f"══════════════════════════════════════{c.END}\n")

    wb.close()


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print(f"\n{c.WARN}Dibatalkan oleh user.{c.END}")
        sys.exit(0)
    except Exception as e:
        print(f"\n{c.ERR}Error: {e}{c.END}")
        sys.exit(1)
